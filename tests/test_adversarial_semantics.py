# Adversarial: the ways the review loop and semantic layer can be wrong
"""Ugly-data and hostile-sequence tests for the semantic and decision layers.

Each test names a way the system could look correct while being wrong, then
proves it is not. Grouped by the surface being attacked.
"""
import time

import pytest
from openpyxl import Workbook
from sqlalchemy import create_engine, event, exc, select
from sqlalchemy.orm import sessionmaker

from backend.app.core.decision_index import DecisionIndex
from backend.app.core.persistence import persist_batch
from backend.app.models.models import (
    Base, DecisionScope, FieldObservation, ProcessingJob, Record,
    ReviewDecision, SourceFile,
)
from engine import ENGINE_VERSION, semantics
from engine.detection import UnreadableFile
from engine.processor import Processor


@pytest.fixture
def db(tmp_path):
    engine = create_engine(f"sqlite:///{tmp_path/'adv.db'}")

    @event.listens_for(engine, "connect")
    def _fk(c, _):
        cur = c.cursor(); cur.execute("PRAGMA foreign_keys=ON"); cur.close()

    Base.metadata.create_all(engine)
    s = sessionmaker(bind=engine)()
    yield s
    s.close()


@pytest.fixture
def job(db):
    src = SourceFile(filename="x.xlsx", stored_path="/tmp/x", size_bytes=1, content_sha256="z")
    db.add(src); db.flush()
    j = ProcessingJob(source_file_id=src.id, status="PROCESSING")
    db.add(j); db.commit()
    return j


def _wb(tmp_path, name, sheets):
    """sheets: {title: [rows]}"""
    wb = Workbook()
    first = True
    for title, rows in sheets.items():
        ws = wb.active if first else wb.create_sheet()
        ws.title = title
        first = False
        for r in rows:
            ws.append(r)
    p = tmp_path / name
    wb.save(p)
    return p


def _ingest(db, job, path, decisions=None):
    def on_batch(rows):
        n = persist_batch(db, rows, job.id); db.commit(); return n
    return Processor(batch_size=100, enable_enrichment=False, decisions=decisions
                     ).process(path, source_name=path.name, on_batch=on_batch)


def _obs(db, field, **where):
    q = select(FieldObservation).where(FieldObservation.canonical_field == field)
    for k, v in where.items():
        q = q.where(getattr(FieldObservation, k) == v)
    return list(db.scalars(q))


def _decide(db, **kw):
    kw.setdefault("scope", DecisionScope.WORKBOOK)
    kw.setdefault("engine_version", ENGINE_VERSION)   # tests may pass an older one
    d = ReviewDecision(decided_by_email="r@example.com", **kw)
    db.add(d); db.commit()
    return d


# ==========================================================================
# 1. The same header meaning different things in different places
# ==========================================================================

def test_same_header_resolves_differently_across_sheets_of_one_workbook(db, job, tmp_path):
    """AREA beside a COMMUNITY column on one sheet, alone on the next."""
    p = _wb(tmp_path, "mixed.xlsx", {
        "WithCommunity": [["NAME", "COMMUNITY", "AREA", "UNIT NO", "MOBILE"],
                          ["A", "Business Bay", "Executive Towers", "1", "+971501111111"]],
        "Alone":         [["NAME", "AREA", "UNIT NO", "MOBILE"],
                          ["B", "Business Bay", "2", "+971502222222"]],
    })
    _ingest(db, job, p)
    with_c = _obs(db, "AREA", source_sheet="WithCommunity")
    alone = _obs(db, "AREA", source_sheet="Alone")
    assert with_c and with_c[0].semantic_type == "sub_community"
    # Correct target, but every signal here is structural -- a neighbouring
    # Community column, a companion header, a vocabulary hit -- and they sum
    # to 0.61. The source never SAID sub-community, so it is flagged: a human
    # confirms it once per workbook and the decision takes over from there.
    assert with_c[0].needs_review and 0.5 < with_c[0].confidence < semantics.min_confidence("AREA")
    # Alone, the only signal is that "Business Bay" is in the known-community
    # vocabulary: a lead, not proof. It reads as community at low confidence
    # and goes to the queue -- never as a confident answer.
    assert alone and alone[0].semantic_type == "community"
    assert alone[0].needs_review and alone[0].confidence < semantics.min_confidence("AREA")
    # and each reached the right canonical column
    a = db.scalar(select(Record).where(Record.unit_number == "1"))
    b = db.scalar(select(Record).where(Record.unit_number == "2"))
    assert a.sub_community == "Executive Towers"
    assert b.community == "Business Bay" and b.sub_community is None


def test_same_header_resolves_differently_across_workbooks(db, job, tmp_path):
    p1 = _wb(tmp_path, "one.xlsx", {"S": [["NAME", "COMMUNITY", "AREA", "UNIT NO", "MOBILE"],
                                          ["A", "JLT", "Cluster D", "1", "+971501111111"]]})
    p2 = _wb(tmp_path, "two.xlsx", {"S": [["NAME", "AREA", "UNIT NO", "MOBILE"],
                                          ["B", "JLT", "2", "+971502222222"]]})
    _ingest(db, job, p1); _ingest(db, job, p2)
    one = _obs(db, "AREA", source_file="one.xlsx")[0]
    two = _obs(db, "AREA", source_file="two.xlsx")[0]
    # Both flagged (structural evidence never clears the threshold on its own),
    # but they resolve to DIFFERENT levels from the same header -- which is the
    # point: no global AREA rule could get both right.
    assert one.semantic_type == "sub_community" and one.needs_review
    assert two.semantic_type == "community" and two.needs_review
    assert one.confidence > two.confidence, "a sibling Community column is stronger evidence than a vocabulary hit"


# ==========================================================================
# 2. Decision precedence under hostile combinations
# ==========================================================================

AMB = {"S1": [["NAME", "UNIT NO", "MOBILE", "Date"], ["A", "1", "+971501111111", "2024-01-01"]],
       "S2": [["NAME", "UNIT NO", "MOBILE", "Date"], ["B", "2", "+971502222222", "2024-02-02"]]}


def test_sheet_scope_applies_to_that_sheet_only(db, job, tmp_path):
    p = _wb(tmp_path, "d.xlsx", AMB)
    _decide(db, canonical_field="Date", original_header="Date", semantic_type="lease_end_date",
            scope=DecisionScope.SHEET, scope_file="d.xlsx", scope_sheet="S1")
    _ingest(db, job, p, decisions=DecisionIndex(db))
    assert _obs(db, "Date", source_sheet="S1")[0].semantic_type == "lease_end_date"
    assert _obs(db, "Date", source_sheet="S2")[0].semantic_type == "unresolved"


def test_workbook_scope_applies_to_every_sheet_of_that_file_and_no_other(db, job, tmp_path):
    p = _wb(tmp_path, "d.xlsx", AMB)
    other = _wb(tmp_path, "e.xlsx", {"S1": AMB["S1"]})
    _decide(db, canonical_field="Date", original_header="Date",
            semantic_type="registration_date", scope_file="d.xlsx")
    idx = DecisionIndex(db)
    _ingest(db, job, p, decisions=idx); _ingest(db, job, other, decisions=idx)
    assert {o.semantic_type for o in _obs(db, "Date", source_file="d.xlsx")} == {"registration_date"}
    assert {o.semantic_type for o in _obs(db, "Date", source_file="e.xlsx")} == {"unresolved"}


def test_global_scope_does_not_override_a_narrower_decision(db, job, tmp_path):
    p = _wb(tmp_path, "d.xlsx", AMB)
    _decide(db, canonical_field="Date", original_header="Date",
            semantic_type="listing_date", scope=DecisionScope.GLOBAL, scope_file=None)
    _decide(db, canonical_field="Date", original_header="Date",
            semantic_type="handover_date", scope_file="d.xlsx")
    _ingest(db, job, p, decisions=DecisionIndex(db))
    assert {o.semantic_type for o in _obs(db, "Date")} == {"handover_date"}


def test_conflicting_decisions_at_the_same_scope_the_newer_one_wins(db, job, tmp_path):
    p = _wb(tmp_path, "d.xlsx", AMB)
    old = _decide(db, canonical_field="Date", original_header="Date",
                  semantic_type="listing_date", scope_file="d.xlsx")
    time.sleep(0.01)
    new = _decide(db, canonical_field="Date", original_header="Date",
                  semantic_type="transaction_date", scope_file="d.xlsx")
    assert new.decided_at > old.decided_at
    _ingest(db, job, p, decisions=DecisionIndex(db))
    assert {o.semantic_type for o in _obs(db, "Date")} == {"transaction_date"}


def test_a_workbook_decision_cannot_be_stored_without_saying_which_workbook(db):
    """Structural: the DB refuses, whatever code tries to write it."""
    db.add(ReviewDecision(canonical_field="Date", semantic_type="handover_date",
                          scope=DecisionScope.WORKBOOK, scope_file=None))
    with pytest.raises(exc.IntegrityError):
        db.commit()
    db.rollback()


def test_a_sheet_decision_cannot_be_stored_without_saying_which_sheet(db):
    db.add(ReviewDecision(canonical_field="Date", semantic_type="handover_date",
                          scope=DecisionScope.SHEET, scope_file="d.xlsx", scope_sheet=None))
    with pytest.raises(exc.IntegrityError):
        db.commit()
    db.rollback()


def test_an_unknown_scope_cannot_be_stored(db):
    db.add(ReviewDecision(canonical_field="Date", semantic_type="handover_date",
                          scope="everywhere", scope_file=None))
    with pytest.raises(exc.IntegrityError):
        db.commit()
    db.rollback()


# ==========================================================================
# 3. Stale decisions, contradiction, immutability
# ==========================================================================

def test_a_decision_from_an_older_engine_is_applied_but_marked_stale(db, job, tmp_path):
    p = _wb(tmp_path, "d.xlsx", {"S1": AMB["S1"]})
    _decide(db, canonical_field="Date", original_header="Date",
            semantic_type="handover_date", scope_file="d.xlsx",
            engine_version=ENGINE_VERSION - 1)
    _ingest(db, job, p, decisions=DecisionIndex(db))
    o = _obs(db, "Date")[0]
    assert o.semantic_type == "handover_date"
    assert o.evidence["decision_stale"] is True
    assert o.evidence["decision_engine_version"] == ENGINE_VERSION - 1


def test_strong_contradictory_evidence_reopens_a_decided_column(db, job, tmp_path):
    """A human decision must not suppress a contradiction forever.

    The decision says lease_start_date. The source says "Handover Date" beside
    a Project column -- strong evidence for handover_date on its own. The
    decision still wins the reading, but the column goes back to the queue
    with the disagreement stated.
    """
    p = _wb(tmp_path, "d.xlsx", {"S": [["NAME", "UNIT NO", "MOBILE", "Handover Date", "Project"],
                                       ["A", "1", "+971501111111", "2023-11-01", "Dubai Hills"]]})
    _decide(db, canonical_field="Date", original_header="Handover Date",
            semantic_type="lease_start_date", scope_file="d.xlsx")
    _ingest(db, job, p, decisions=DecisionIndex(db))
    o = _obs(db, "Date")[0]
    assert o.semantic_type == "lease_start_date", "decision must still win the reading"
    assert o.needs_review is True, "contradiction was suppressed"
    assert o.evidence["contradicts_inference"] is True
    assert o.evidence["inferred_semantic_type"] == "handover_date"
    assert "contradicted" in o.evidence["reason"]


def test_weak_contradictory_evidence_does_not_reopen_a_decided_column(db, job, tmp_path):
    """The mirror: a bare 'Date' column infers nothing strong, so a decision
    about it stands without being re-questioned on every ingest."""
    p = _wb(tmp_path, "d.xlsx", {"S1": AMB["S1"]})
    _decide(db, canonical_field="Date", original_header="Date",
            semantic_type="transaction_date", scope_file="d.xlsx")
    _ingest(db, job, p, decisions=DecisionIndex(db))
    o = _obs(db, "Date")[0]
    assert o.semantic_type == "transaction_date" and o.needs_review is False


def test_recording_a_decision_does_not_touch_rows_already_ingested(db, job, tmp_path):
    p = _wb(tmp_path, "d.xlsx", {"S1": AMB["S1"]})
    _ingest(db, job, p)
    before = [(o.id, o.semantic_type, o.needs_review, o.evidence) for o in _obs(db, "Date")]
    _decide(db, canonical_field="Date", original_header="Date",
            semantic_type="handover_date", scope_file="d.xlsx")
    db.expire_all()
    after = [(o.id, o.semantic_type, o.needs_review, o.evidence) for o in _obs(db, "Date")]
    assert before == after, "a decision mutated historical observations"


def test_reprocessing_is_what_applies_a_decision_to_existing_rows(db, job, tmp_path):
    """Same file, ingested again with the decision present = the reprocess path."""
    p = _wb(tmp_path, "d.xlsx", {"S1": AMB["S1"]})
    _ingest(db, job, p)
    assert _obs(db, "Date")[0].semantic_type == "unresolved"
    _decide(db, canonical_field="Date", original_header="Date",
            semantic_type="handover_date", scope_file="d.xlsx")
    for o in _obs(db, "Date"):
        db.delete(o)
    db.commit()
    _ingest(db, job, p, decisions=DecisionIndex(db))
    assert _obs(db, "Date")[0].semantic_type == "handover_date"


def test_a_degraded_decision_index_is_reported_on_the_job(db, job, tmp_path):
    class Degraded:
        degraded_reason = "OperationalError: no such table: review_decisions"
        def lookup(self, *a, **k): return None
    res = _ingest(db, job, _wb(tmp_path, "d.xlsx", {"S1": AMB["S1"]}), decisions=Degraded())
    codes = [e["code"] for e in res.errors]
    assert "DECISIONS_UNAVAILABLE" in codes


def test_a_failing_lookup_marks_the_reading_inference_only(db, job, tmp_path):
    class Exploding:
        def lookup(self, *a, **k): raise RuntimeError("boom")
    _ingest(db, job, _wb(tmp_path, "d.xlsx", {"S1": AMB["S1"]}), decisions=Exploding())
    o = _obs(db, "Date")[0]
    assert o.needs_review is True
    assert "decision_lookup_failed" in o.evidence


# ==========================================================================
# 4. Hostile files
# ==========================================================================

def test_a_missing_source_file_raises_rather_than_producing_an_empty_success(db, job, tmp_path):
    with pytest.raises((UnreadableFile, FileNotFoundError, OSError)):
        _ingest(db, job, tmp_path / "does_not_exist.xlsx")


def test_a_malformed_workbook_raises_rather_than_producing_an_empty_success(db, job, tmp_path):
    p = tmp_path / "garbage.xlsx"
    p.write_bytes(b"this is not a zip container at all " * 50)
    with pytest.raises(UnreadableFile):
        _ingest(db, job, p)


def test_hidden_columns_are_still_ingested(db, job, tmp_path):
    wb = Workbook(); ws = wb.active; ws.title = "S"
    ws.append(["NAME", "UNIT NO", "MOBILE", "Handover Date"])
    ws.append(["A", "1", "+971501111111", "2023-11-01"])
    ws.column_dimensions["D"].hidden = True
    p = tmp_path / "hidden.xlsx"; wb.save(p)
    _ingest(db, job, p)
    assert _obs(db, "Date"), "a hidden column's data was dropped"


def test_a_merged_header_cell_does_not_crash_the_sheet(db, job, tmp_path):
    wb = Workbook(); ws = wb.active; ws.title = "S"
    ws.append(["NAME", "CONTACT", None, "UNIT NO"])
    ws.merge_cells("B1:C1")
    ws.append(["A", "+971501111111", "a@example.com", "1"])
    p = tmp_path / "merged.xlsx"; wb.save(p)
    res = _ingest(db, job, p)
    assert not [e for e in res.errors if e["code"] == "SHEET_FAILED"]
    assert db.scalar(select(Record).limit(1)) is not None


def test_an_invalid_numeric_size_keeps_its_raw_value_and_parses_to_nothing(db, job, tmp_path):
    p = _wb(tmp_path, "s.xlsx", {"S": [["NAME", "UNIT NO", "MOBILE", "Size (sq.ft)"],
                                       ["A", "1", "+971501111111", "twelve hundred"]]})
    _ingest(db, job, p)
    o = _obs(db, "Size")[0]
    assert o.raw_value == "twelve hundred"
    assert o.parsed_number is None
    rec = db.scalar(select(Record))
    assert rec.size is None, "an unparseable size became a number"


def test_formula_cells_do_not_crash_and_preserve_what_was_read(db, job, tmp_path):
    wb = Workbook(); ws = wb.active; ws.title = "S"
    ws.append(["NAME", "UNIT NO", "MOBILE", "Size (sq.ft)"])
    ws.append(["A", "1", "+971501111111", "=1000+200"])
    p = tmp_path / "formula.xlsx"; wb.save(p)
    res = _ingest(db, job, p)
    assert not [e for e in res.errors if e["code"] == "SHEET_FAILED"]
    o = _obs(db, "Size")
    # openpyxl without a calc engine yields either the formula text or None;
    # either way the raw value must be what was read and nothing invented.
    if o:
        assert o[0].raw_value in ("=1000+200", None) or o[0].parsed_number in (None, 1200.0)


def test_blank_and_whitespace_only_rows_are_skipped_not_stored(db, job, tmp_path):
    p = _wb(tmp_path, "b.xlsx", {"S": [["NAME", "UNIT NO", "MOBILE"],
                                       ["A", "1", "+971501111111"],
                                       ["", "", ""], [" ", None, ""],
                                       ["B", "2", "+971502222222"]]})
    _ingest(db, job, p)
    assert db.scalar(select(Record).where(Record.name == "A")) is not None
    assert len(list(db.scalars(select(Record)))) == 2
